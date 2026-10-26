r"""List every part number in the vault and the drawing that shows it, from a DmProbe of P:\APEX (read-only).

One row per PartNo (fasteners, FAST-, left out): description, category (last letter of the code), group (0806,
0811 incl. 08Z11, Shared = 08BXX / 08XXX / GNRIC), place tag, file, configuration, vault folder, PartNoStatus,
drawing, note. The drawing is the one named for the PartNo (D2), else any drawing whose views show that file and
configuration (a view of an As Welded child counts for its As Machined parent; a file-level number counts any view
of the file). GNRIC-LMP and GNRIC-PKP are noted as likely cut lists. A Summary sheet counts per group.
Re-run after each import or when drawings are added: probe P:\APEX first (DmProbe P:\APEX <probeDir>).

Usage: python -I drawing_list.py <probeDir> <out.xlsx>
"""
import collections, csv, os, re, sys
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

if len(sys.argv) != 3:
    sys.exit("Usage: python -I drawing_list.py <probeDir> <out.xlsx>")
probe, out = sys.argv[1], sys.argv[2]
rd = lambda n: list(csv.DictReader(open(os.path.join(probe, n), encoding="utf-8-sig")))
files, configs, props, views = rd("files.csv"), rd("configs.csv"), rd("properties.csv"), rd("views.csv")

folder = {f["file"].lower(): f["folder"] for f in files}
drawings = {f["file"][:-7].upper(): f["file"] for f in files if f["file"].lower().endswith(".slddrw")}
prop = {(p["file"].lower(), p["config"], p["name"].lower()): p["value"] for p in props}
parent = {(c["file"].lower(), c["config"]): c["parent_config"] for c in configs}

# Which drawings show each file / configuration (a view of an As Welded child counts for its As Machined parent).
shown = collections.defaultdict(set)
for v in views:
    m, c = v["model"].lower(), v["config"]
    shown[(m, c)].add(v["drawing"])
    shown[(m, "*")].add(v["drawing"])
    if parent.get((m, c)):
        shown[(m, parent[(m, c)])].add(v["drawing"])

rows = []
for c in configs:
    if c["PartNo"]:
        rows.append((c["PartNo"], c["file"], c["config"]))
for f in files:
    if f["file_PartNo"]:
        rows.append((f["file_PartNo"], f["file"], "(file)"))

def group(pn):
    t = pn.split("-")[0].upper()
    if t in ("GNRIC", "08BXX", "08XXX"):
        return "Shared"
    if re.match(r"^08[A-Z]06$", t):
        return "0806"
    if re.match(r"^08[A-Z]11$", t):
        return "0811"
    return "?"

table, seen = [], collections.Counter()
for pn, f, cfg in rows:
    if pn.upper().startswith("FAST-"):
        continue
    fk = f.lower()
    level = "(file)" if cfg == "(file)" else cfg
    desc = prop.get((fk, cfg, "description")) or prop.get((fk, "(file)", "description")) or ""
    status = prop.get((fk, level, "partnostatus")) or (prop.get((fk, "(file)", "partnostatus")) if cfg != "(file)" else "") or ""
    code = pn.split("-")[1] if pn.count("-") >= 2 else ""
    dw = drawings.get(pn.upper())                                   # D2: drawing named for the PartNo
    if not dw:
        hits = sorted(shown.get((fk, "*" if cfg == "(file)" else cfg), ()))
        dw = "; ".join(hits)
    note = "likely cut list" if re.match(r"^GNRIC-(LMP|PKP)-", pn, re.I) else ""
    seen[pn.upper()] += 1
    table.append([pn, desc, code[-1:] if code else "", group(pn), pn.split("-")[0], f, cfg, "APEX\\" + folder.get(fk, ""), status, dw, note])
table.sort(key=lambda r: ({"0806": 0, "0811": 1, "Shared": 2}.get(r[3], 3), r[0], r[6]))
dupes = {k for k, n in seen.items() if n > 1}

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Part numbers"
head = ["PartNo", "Description", "Category", "Group", "Place tag", "File", "Configuration", "Vault folder", "PartNoStatus", "Drawing", "Note"]
ws.append(head)
for r in table:
    if r[0].upper() in dupes:
        r[10] = "; ".join(x for x in (r[10], "PartNo on more than one file or configuration") if x)
    ws.append(r)

def tidy(ws):
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 60)
tidy(ws)

s = wb.create_sheet("Summary")
s.append(["Group", "Part numbers", "With a drawing", "Without a drawing", "Without, not counting likely cut lists"])
tot = [0, 0, 0, 0]
for g in ("0806", "0811", "Shared"):
    rs = [r for r in table if r[3] == g]
    n, w = len(rs), sum(1 for r in rs if r[9])
    wo = n - w
    wo_nc = sum(1 for r in rs if not r[9] and not r[10].startswith("likely cut list"))
    s.append([g, n, w, wo, wo_nc])
    tot = [a + b for a, b in zip(tot, (n, w, wo, wo_nc))]
s.append(["Total"] + tot)
for c in s[s.max_row]:
    c.font = Font(bold=True)
tidy(s)
wb.save(out)

print("rows", len(table), "dupes", sorted(dupes), "groups", collections.Counter(r[3] for r in table))
print("with drawing", sum(1 for r in table if r[9]), "drawing files in vault", len(drawings))
used = set(d for r in table for d in r[9].split("; ") if d)
print("drawings not matched to any row:", sorted(set(drawings.values()) - used))
for r in s.iter_rows(values_only=True):
    print(" ", r)
