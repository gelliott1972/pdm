r"""J2 / J3 property sheet: Description, Finish and Weight for every vault part number, and every DrawingRef to delete.
Read-only: works from a DmProbe of P:\APEX and P:\_Shared; writes one .xlsx for Grant and James to review.

Agreed rules (Admin, 28 Oct), following James's 08A11-SFA:
  Finish once at file level; Description and Weight per configuration, Weight = "SW-Mass@@<config>@<file>".
  A file whose PartNo is at file level (0806 precedent) keeps file level: Weight = "SW-Mass@<file>".
  Every DrawingRef (file and configuration level) is deleted. Drawings are left out. As Welded configurations carry no
  PartNo: listed on their own sheet, not written. D17: the 9 fastener configurations that get a PartNo in the same run.
Proposals: Weight from the formula; Finish from the same file, else its 0811 / 0806 equivalent, else vault files of the same
category and type, e.g. all HBA (marked);
Description from the row's own value, else its 0811 / 0806 equivalent (same number, other screen; marked), else blank.

Sheets: Properties (one row per file + configuration with a PartNo; the writer reads the Proposed columns),
DrawingRef (all of them, deleted by the writer), As Welded (information only), Counts.

With --since <oldProbeDir> (the probe an earlier sheet was built from; Admin, 29 Oct): a value someone changed in the
vault since then is taken as current and never proposed for overwriting. Rows with such a value are marked in a
"Changed since" column, and a "Changed since" sheet lists every change, flagging formulas that point at another file
("needs fixing by James") and file-level Description / Weight on files whose PartNo is per configuration (left as is
until Grant decides).

Usage: python -I props_sheet.py <probeDir> <ledgerCsv> <out.xlsx> [--since <oldProbeDir>]
       (probeDir holds APEX\ and _Shared\ from DmProbe)
"""
import collections, csv, os, re, sys
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

args = sys.argv[1:]
since = None
if "--since" in args:
    i = args.index("--since")
    since = args[i + 1] if i + 1 < len(args) else None
    del args[i:i + 2]
if len(args) != 3 or ("--since" in sys.argv and not since):
    sys.exit("Usage: python -I props_sheet.py <probeDir> <ledgerCsv> <out.xlsx> [--since <oldProbeDir>]")
probe, ledger_path, out = args

# D17 (agreed 25 Oct): fastener configurations used by the vault that get a PartNo by check-out.
D17 = {
    ("FAST-HEX-8.8-HDG.SLDPRT", "ISO 4017 - M10 x 140-N"): "FAST-HEX-M10x140-8.8-HDG",
    ("FAST-HEX-8.8-HDG.SLDPRT", "ISO 4017 - M10 x 60-N"): "FAST-HEX-M10x60-8.8-HDG",
    ("FAST-HEX-8.8-HDG.SLDPRT", "ISO 4017 - M6 x 50-N"): "FAST-HEX-M6x50-8.8-HDG",
    ("FAST-HEX-8.8-HDG.SLDPRT", "ISO 4017 - M8 x 55-N"): "FAST-HEX-M8x55-8.8-HDG",
    ("FAST-HEX-8.8-HDG.SLDPRT", "ISO 4017 - M8 x 90-N"): "FAST-HEX-M8x90-8.8-HDG",
    ("FAST-NUT-8.8-HDG.SLDPRT", "BS EN 24032 - M6 - W - N"): "FAST-NUT-M6-8.8-HDG",
    ("FAST-NUT-8.8-HDG.SLDPRT", "BS EN 24032 - M8 - W - N"): "FAST-NUT-M8-8.8-HDG",
    ("FAST-WSH-8.8-HDG.SLDPRT", "Washer ISO 7089 - 6"): "FAST-WSH-M6-8.8-HDG",
    ("FAST-WSH-8.8-HDG.SLDPRT", "Washer ISO 7089 - 8"): "FAST-WSH-M8-8.8-HDG",
}
CATEGORY = {"HB": "Header Beam", "SF": "Subframe", "LM": "LED Member", "LP": "LED Plate", "PL": "Standard Plate",
            "BR": "Brace", "DB": "Diagonal Brace", "BK": "Bracket", "PK": "Packer / Shim", "FS": "Fastener Kit",
            "MS": "Miscellaneous Steel", "ST": "Structure"}
TYPE = {"A": "assembly", "W": "weldment", "P": "part"}

files, configs, props = [], [], []
for area in ("APEX", "_Shared"):
    d = os.path.join(probe, area)
    rd = lambda n: list(csv.DictReader(open(os.path.join(d, n), encoding="utf-8-sig")))
    for f in rd("files.csv"):
        f["area"] = area
        files.append(f)
    configs += rd("configs.csv")
    props += rd("properties.csv")
ledger = list(csv.DictReader(open(ledger_path, encoding="utf-8-sig")))

folder = {f["file"].lower(): f["area"] + "\\" + f["folder"] for f in files}
expr = {}                                            # (file, config or "(file)", name) -> expression as stored
for p in props:
    if p["name"]:
        expr[(p["file"].lower(), p["config"], p["name"].lower())] = p["expression"] or p["value"]
get = lambda f, c, n: (expr.get((f.lower(), c, n.lower())) or "").strip()

def strip_q(s):
    return s[1:-1] if len(s) >= 2 and s[0] == s[-1] == '"' else s

# Changes since an earlier probe: (file, config or "(file)", name) -> (old, new); None = not there.
WATCH = ("description", "finish", "weight", "material", "partno", "drawingref")
changed = {}
if since:
    old = {}
    for area in ("APEX", "_Shared"):
        for p in csv.DictReader(open(os.path.join(since, area, "properties.csv"), encoding="utf-8-sig")):
            if p["name"]:
                old[(p["file"].lower(), p["config"], p["name"].lower())] = p["expression"] or p["value"]
    for k in set(old) | set(expr):
        if k[2] in WATCH and old.get(k) != expr.get(k):
            changed[k] = (old.get(k), expr.get(k))

def other_file(f, v):
    """A SW-xxx@ formula that names a file other than f (copied from another file): the file it names."""
    m = re.search(r'SW-[\w ]+@(?:@[^@"]*@)?([^@"]+\.SLD(?:PRT|ASM))', v or "", re.I)
    return m.group(1) if m and m.group(1).lower() != f.lower() else None

def weight_formula(f, cfg):
    return '"SW-Mass@%s"' % f if cfg == "(file)" else '"SW-Mass@@%s@%s"' % (cfg, f)

# Old names (where the file came from), for the reviewer.
old_name = {}
for r in ledger:
    k = (r["new_file"].lower(), r["new_config"] or "")
    if k not in old_name:
        old_name[k] = r["old_file"] + (" / " + r["old_config"] if r["old_config"] else "")

# Rows: every non-drawing file + configuration with a PartNo (file-level PartNo -> one "(file)" row), plus D17.
rows, welded = [], []
pn_of = {}
for c in configs:
    pn_of[(c["file"].lower(), c["config"])] = c["PartNo"]
for f in files:
    fn = f["file"]
    if fn.lower().endswith(".slddrw"):
        continue
    if f["file_PartNo"]:
        rows.append(dict(file=fn, cfg="(file)", pn=f["file_PartNo"], d17=False))
    for c in configs:
        if c["file"].lower() != fn.lower():
            continue
        if c["PartNo"]:
            rows.append(dict(file=fn, cfg=c["config"], pn=c["PartNo"], d17=False))
        elif (fn, c["config"]) in D17:
            rows.append(dict(file=fn, cfg=c["config"], pn=D17[(fn, c["config"])], d17=True))
        elif "<as welded>" in c["config"].lower():
            welded.append(c)
missing17 = set(D17) - set((r["file"], r["cfg"]) for r in rows if r["d17"])
if missing17:
    sys.exit("D17 configurations not in the probe: %s" % sorted(missing17))

def group(pn):
    t = pn.split("-")[0].upper()
    if t == "FAST":
        return "Fasteners"
    if t in ("GNRIC", "08BXX", "08XXX"):
        return "Shared"
    m = re.match(r"^08[A-Z](06|11)$", t)
    return "08" + m.group(1) if m else "?"

def code(pn):
    parts = pn.split("-")
    return parts[1].upper() if len(parts) >= 3 else ""

# Current values: configuration level first, then file level (marked).
def current(f, cfg, name):
    if cfg != "(file)":
        v = get(f, cfg, name)
        if v:
            return v, "config"
    v = get(f, "(file)", name)
    return (v, "file") if v else ("", "")

# Finish of each part number and by category + type (HBA, SFA, LMW...) among vault files that have one.
finish_by_pn, finish_by_cat = {}, collections.defaultdict(collections.Counter)
for r in rows:
    if r["d17"] or group(r["pn"]) == "Fasteners":
        continue
    v = get(r["file"], "(file)", "finish") or (get(r["file"], r["cfg"], "finish") if r["cfg"] != "(file)" else "")
    if v and v != ".":
        finish_by_pn[r["pn"].upper()] = v
        finish_by_cat[code(r["pn"])][v] += 1

# Descriptions by PartNo, for the 0811 / 0806 equivalent.
desc_by_pn = {}
for r in rows:
    v, _ = current(r["file"], r["cfg"], "description")
    if v:
        desc_by_pn[r["pn"].upper()] = v

def equivalent(pn, known):
    """Same category and number on the other screen: 08A06-SFA-1T <-> 08A11-SFA-1T (B/C also via 08Z11)."""
    t, rest = pn.split("-", 1)
    m = re.match(r"^08([A-Z])(06|11)$", t.upper())
    if not m:
        return None
    shaft, scr = m.groups()
    other = "11" if scr == "06" else "06"
    tags = ["08%s%s" % (shaft, other)]
    if other == "11" and shaft in "BC":
        tags.append("08Z11")
    if shaft == "Z":
        tags = ["08B" + other, "08C" + other]
    for tg in tags:
        k = (tg + "-" + rest).upper()
        if k in known:
            return k
    return None

def screen_name(pn):
    """08A11-... -> LS-08A-11 (the screen name James uses in descriptions)."""
    t = pn.split("-")[0].upper()
    return "LS-%s-%s" % (t[:3], t[3:])

pn_row = {r["pn"].upper(): r for r in rows}
def came_from(pn):
    r = pn_row[pn]
    return old_name.get((r["file"].lower(), r["cfg"] if r["cfg"] != "(file)" else ""), "") or old_name.get((r["file"].lower(), ""), "") or r["file"]

table = []
for r in rows:
    f, cfg, pn = r["file"], r["cfg"], r["pn"]
    g = group(pn)
    cd, cd_lvl = current(f, cfg, "description")
    cfin, cfin_lvl = current(f, cfg, "finish")
    cw, cw_lvl = current(f, cfg, "weight")
    dref = [x for x in (("file", get(f, "(file)", "drawingref")), ("config", get(f, cfg, "drawingref") if cfg != "(file)" else "")) if x[1]]
    notes, guess = [], []

    # Weight: the formula for this row's level.
    pw = weight_formula(f, cfg)
    # Description: own value (moved to configuration level if it sits at file level), else the equivalent.
    if cd:
        pd = cd
        if cd_lvl == "file" and cfg != "(file)":
            notes.append("Description copied from file level to the configuration")
    elif not r["d17"] and equivalent(pn, desc_by_pn):
        eq = equivalent(pn, desc_by_pn)
        pd = re.sub(screen_name(eq), screen_name(pn), desc_by_pn[eq])
        guess.append("Description from %s, was %s (check)" % (eq, came_from(eq)) + (", screen name changed" if pd != desc_by_pn[eq] else ""))
    else:
        pd = ""
        guess.append("Description BLANK")
    if pd.upper().endswith("ASSYY"):
        notes.append("Description ends ASSYY: typo?")
    if r["d17"] or g == "Fasteners":
        m = re.search(r"M(\d+) x (\d+)", cfg)
        if m and pd and ("M%sx%s" % m.groups()) not in pd.replace(" ", ""):
            notes.append("Description says %s, configuration is M%sx%s: correct it?" % (pd, m.group(1), m.group(2)))
    # Finish: once at file level.
    ffile = get(f, "(file)", "finish")
    fcfg = sorted(set(get(f, c["config"], "finish") for c in configs if c["file"].lower() == f.lower()) - {""})
    if ffile and ffile != ".":
        pf = ffile
    elif len(fcfg) == 1:
        pf = fcfg[0]
        notes.append("Finish moved from configuration to file level")
    elif len(fcfg) > 1:
        pf = ""
        guess.append("Finish BLANK: configurations differ (%s)" % ", ".join(fcfg))
    elif (g == "Fasteners" or r["d17"]) and "-HDG." in f.upper():
        pf = "HDG"
        guess.append("Finish HDG from the file name (check)")
    elif equivalent(pn, finish_by_pn):
        eq = equivalent(pn, finish_by_pn)
        pf = finish_by_pn[eq]
        guess.append("Finish %s from %s, was %s (check)" % (pf, eq, came_from(eq)))
    else:
        common = finish_by_cat.get(code(pn))
        if common and len(common) == 1:
            pf = next(iter(common))
            guess.append("Finish %s as other %s files (check)" % (pf, code(pn)))
        else:
            pf = ""
            guess.append("Finish BLANK" + (": %s files differ (%s)" % (code(pn), ", ".join(sorted(common))) if common else ""))
    if ffile == ".":
        notes.append('file-level Finish is "."')
    if cw and strip_q(cw) != strip_q(pw):
        notes.append("Weight formula changes from %s" % cw)
    if r["d17"]:
        notes.append("D17: PartNo added in the same run (no PartNoStatus, as the other fasteners)")

    # Values changed since the earlier probe stay as they are now (never proposed for overwriting).
    chg = sorted(k for k in changed if k[0] == f.lower() and k[1] in (cfg, "(file)"))
    lvl_now = lambda n, c: get(f, c, n)
    if (f.lower(), cfg, "description") in changed:
        pd = lvl_now("description", cfg)
    if (f.lower(), "(file)", "finish") in changed:
        pf = ffile
    if (f.lower(), cfg, "weight") in changed:
        pw = lvl_now("weight", cfg)
    since_txt = []
    for k in chg:
        o, n = changed[k]
        t = "%s [%s]: %s -> %s" % (k[2], "file" if k[1] == "(file)" else "config", "(none)" if o is None else o or "(empty)",
                                   "(none)" if n is None else n or "(empty)")
        if other_file(f, n):
            t += "  NEEDS FIXING BY JAMES: points at %s" % other_file(f, n)
        since_txt.append(t)

    nothing = (pd == cd and cd_lvl in ("config", "") if cfg != "(file)" else pd == cd) and \
              (pf == ffile) and (strip_q(cw) == strip_q(pw) and (cw_lvl == "config" or cfg == "(file)"))
    if r["d17"]:
        nothing = False
    blank = not pd or not pf
    status = "blank" if blank else ("complete" if nothing else "pre-filled")
    hint = old_name.get((f.lower(), cfg if cfg != "(file)" else ""), "") or old_name.get((f.lower(), ""), "")
    cat = code(pn)
    hint2 = "%s %s" % (CATEGORY.get(cat[:2], ""), TYPE.get(cat[2:3], "")) if cat else ""
    lv = lambda v, l: (v + " [file]") if v and l == "file" and cfg != "(file)" else v
    table.append([g, pn, f, cfg, folder.get(f.lower(), ""),
                  lv(cd, cd_lvl), lv(cfin, cfin_lvl), lv(cw, cw_lvl),
                  "; ".join("%s: %s" % x for x in dref),
                  pd, pf, pw, status, "; ".join(guess), "; ".join(notes), hint, hint2.strip()]
                 + (["; ".join(since_txt)] if since else []))

order = {"08_06": 0, "0806": 0, "0811": 1, "Shared": 2, "Fasteners": 3}
table.sort(key=lambda t: (order.get(t[0], 4), t[1], t[3]))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Properties"
head = ["Group", "PartNo", "File", "Configuration", "Vault folder",
        "Current Description", "Current Finish", "Current Weight", "DrawingRef (to delete)",
        "Proposed Description", "Proposed Finish (file level)", "Proposed Weight",
        "Status", "Guesses to check", "Notes", "Came from (ledger)", "Category"] + (["Changed since"] if since else [])
ws.append(head)
yellow = PatternFill("solid", fgColor="FFF2CC")
red = PatternFill("solid", fgColor="F8CBAD")
blue = PatternFill("solid", fgColor="BDD7EE")
orange = PatternFill("solid", fgColor="F4B084")
for t in table:
    ws.append(t)
    row = ws.max_row
    for col, name in ((10, "Description"), (11, "Finish")):
        cell = ws.cell(row=row, column=col)
        if not cell.value:
            cell.fill = red
        elif name in t[13]:
            cell.fill = yellow
    if since and t[17]:
        for col in (2, 18):
            ws.cell(row=row, column=col).fill = orange if "NEEDS FIXING" in t[17] else blue

def tidy(ws, maxw=50):
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), maxw)
tidy(ws)

d = wb.create_sheet("DrawingRef")
d.append(["File", "Level", "Value (to delete)", "Group"])
drefs = [p for p in props if p["name"].lower() == "drawingref"]
for p in sorted(drefs, key=lambda p: (p["file"].lower(), p["config"])):
    pn = pn_of.get((p["file"].lower(), p["config"])) or next((r["pn"] for r in rows if r["file"] == p["file"]), "")
    d.append([p["file"], p["config"], p["expression"] or p["value"], group(pn) if pn else ""])
tidy(d)

w = wb.create_sheet("As Welded")
w.append(["File", "Configuration", "As Machined parent", "PartNo of parent", "Current Description", "Current Finish", "Current Weight", "Not written unless Admin says"])
for c in sorted(welded, key=lambda c: (c["file"].lower(), c["config"])):
    f = c["file"]
    w.append([f, c["config"], c["parent_config"], pn_of.get((f.lower(), c["parent_config"]), ""),
              get(f, c["config"], "description"), get(f, c["config"], "finish"), get(f, c["config"], "weight"), ""])
tidy(w)

if since:
    file_pn = {f["file"].lower() for f in files if f["file_PartNo"]}
    c = wb.create_sheet("Changed since")
    c.append(["File", "Level", "Property", "Before (%s)" % os.path.basename(since.rstrip("\\/")),
              "Now (%s)" % os.path.basename(probe.rstrip("\\/")), "PartNo is", "Flag", "In the Properties sheet"])
    sheet_files = {t[2].lower() for t in table}
    name_of = {f["file"].lower(): f["file"] for f in files}
    for k in sorted(changed):
        o, n = changed[k]
        flag = []
        if other_file(k[0], n):
            flag.append("NEEDS FIXING BY JAMES: formula points at %s" % other_file(k[0], n))
        perconf = k[0] not in file_pn
        if k[1] == "(file)" and k[2] in ("description", "weight") and perconf and n:
            flag.append("file level, but PartNo is per configuration: left as is until Grant decides")
        fn = name_of.get(k[0], k[0])
        c.append([fn, "file" if k[1] == "(file)" else k[1], k[2], "(none)" if o is None else o, "(none)" if n is None else n,
                  "per configuration" if perconf else "file level", "; ".join(flag),
                  "kept as now, not overwritten" if k[0] in sheet_files else "not in the sheet (drawing or no PartNo)"])
        if flag:
            c.cell(row=c.max_row, column=7).fill = orange if "NEEDS" in flag[0] else blue
    tidy(c, 60)

s = wb.create_sheet("Counts")
s.append(["Group", "Rows", "Complete (nothing to write)", "Pre-filled", "Blank (something to fill)", "Pre-filled with a guess to check", "DrawingRef on these rows"])
tot = [0] * 6
for g in ("0806", "0811", "Shared", "Fasteners"):
    rs = [t for t in table if t[0] == g]
    vals = [len(rs), sum(t[12] == "complete" for t in rs), sum(t[12] == "pre-filled" for t in rs),
            sum(t[12] == "blank" for t in rs), sum(t[12] == "pre-filled" and bool(t[13]) for t in rs), sum(bool(t[8]) for t in rs)]
    s.append([g] + vals)
    tot = [a + b for a, b in zip(tot, vals)]
s.append(["Total"] + tot)
s.append([])
s.append(["DrawingRef rows in the vault (all deleted)", len(drefs)])
s.append(["As Welded configurations (listed, not written)", len(welded)])
if since:
    s.append(["Rows with a value changed since %s (kept as now)" % os.path.basename(since.rstrip("\\/")), sum(bool(t[17]) for t in table)])
    s.append(["Property values changed since then (Changed since sheet)", len(changed)])
for c in s[s.max_row - (5 if since else 3)]:
    c.font = Font(bold=True)
tidy(s, 60)
wb.save(out)

for r in s.iter_rows(values_only=True):
    print(" ", r)
print("blank by reason:", collections.Counter(x.split(":")[0] for t in table for x in t[13].split("; ") if "BLANK" in x))
print("guesses:", collections.Counter(re.sub(r" from \S+| as other .*| \(check\)", "", x) for t in table for x in t[13].split("; ") if x and "BLANK" not in x))
print("notes:", collections.Counter(re.sub(r" from .*| \(.*| [A-Z0-9]{5}-.*|Description says.*", "", x) for t in table for x in t[14].split("; ") if x))
